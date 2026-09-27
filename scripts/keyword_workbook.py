"""Build the colour-coded keyword + content-calendar workbook from the seo-research/ CSVs (the source of truth).

Re-run after every publish, once the CSV Status columns are updated:  python3 scripts/keyword_workbook.py
"""
import csv
import os

import openpyxl
from openpyxl.styles import Font, PatternFill

C = "seo-research/Blog_Keywords_and_Content_Calendar/"
SHEETS = [
    ("Content_Calendar", C + "Task8_26Week_Content_Calendar.csv"),
    ("Blog_Keywords_CA", C + "Table_4A_Blog_Keywords_CA.csv"),
    ("Blog_Keywords_US", C + "Table_4A_Blog_Keywords_US.csv"),
    ("Question_Keywords", C + "Table_4B_Question_Keywords_FAQ_vs_Standalone.csv"),
    ("Seasonal", C + "Table_4C_Seasonal_Calendar.csv"),
    ("Master_Keywords_CA", "seo-research/Table_3A_Master_Keyword_List_CA.csv"),
    ("Keyword_Gap", "seo-research/Table_3B_Keyword_Gap_Opportunities.csv"),
    ("Keyword_Clusters", "seo-research/Table_5_Keyword_Clusters.csv"),
]
OUT = os.path.expanduser("~/Documents/All website keywords Data/timhortonsdonuts.com/Tim Hortons Keywords & Content Calendar.xlsx")
# Status prefix -> fill (same meaning across sheets).
FILLS = {
    "Published": "C6EFCE",          # green: page live and owns the keyword
    "Covered": "DDEBF7",            # blue: answered by a section of a live page - no new page
    "Merge-Recommended": "FFC7CE",  # red: planned page would cannibalise a live page
    "Partial": "FFEB9C",            # amber: overlap - only a different angle is safe
    "Planned": "FFEB9C",
}


def fill_for(status):
    return next((PatternFill("solid", fgColor=c) for p, c in FILLS.items() if status.startswith(p)), None)


wb = openpyxl.Workbook()
wb.remove(wb.active)
for name, path in SHEETS:
    rows = list(csv.reader(open(path, newline="")))
    ws = wb.create_sheet(name)
    for r in rows:
        ws.append([int(v.replace(",", "")) if v.replace(",", "").isdigit() else v for v in r])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1A1A1A")
    if "Status" in rows[0]:
        col = rows[0].index("Status") + 1
        for r in range(2, ws.max_row + 1):
            cell = ws.cell(r, col)
            if cell.value and (f := fill_for(cell.value)):
                cell.fill, cell.font = f, Font(bold=True)
    for i, h in enumerate(rows[0], 1):
        width = max(len(str(r[i - 1])) for r in rows if len(r) >= i)
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = min(max(width, len(h)) + 2, 60)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print(f"wrote {OUT} ({len(SHEETS)} sheets)")
